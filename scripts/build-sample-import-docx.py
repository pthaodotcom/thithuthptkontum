from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out=Path(__file__).resolve().parent.parent/"test-data";out.mkdir(exist_ok=True)
doc=Document();section=doc.sections[0];section.orientation=1;section.page_width=Cm(29.7);section.page_height=Cm(21)
section.left_margin=section.right_margin=Cm(1.2);section.top_margin=section.bottom_margin=Cm(1.2)
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run("MẪU IMPORT CÂU HỎI - WORD");r.bold=True;r.font.size=Pt(16);r.font.color.rgb=RGBColor(29,78,216)
doc.add_paragraph("Giữ nguyên tên cột. Mỗi dòng là một câu hỏi; Phần I/II dùng đủ bốn đáp án, Phần III dùng Đáp án đúng gồm đúng bốn ký tự.")
headers=["Phan","ChuyenDe","BaiHoc","MucDo","NoiDung","DapAn1","DapAn2","DapAn3","DapAn4","DapAnDung"]
rows=[
["I","Hàm số","Đạo hàm","Nhận biết","Cho f(x)=x². Tính f'(x).","x","2x","x²","2","2"],
["II","Hàm số","Đạo hàm","Thông hiểu","Xét đúng/sai.","Ý a","Ý b","Ý c","Ý d","Đ;S;Đ;S"],
["III","Hàm số","Đạo hàm","Vận dụng","Điền bốn ký tự.","","","","","-1,2"]]
t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.style="Table Grid"
for i,h in enumerate(headers):
 c=t.rows[0].cells[i];c.text=h;c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
 shd=OxmlElement("w:shd");shd.set(qn("w:fill"),"1D4ED8");c._tc.get_or_add_tcPr().append(shd)
 for run in c.paragraphs[0].runs:run.bold=True;run.font.color.rgb=RGBColor(255,255,255);run.font.size=Pt(8)
for data in rows:
 cells=t.add_row().cells
 for i,v in enumerate(data):
  cells[i].text=v;cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  for run in cells[i].paragraphs[0].runs:run.font.size=Pt(8)
doc.save(out/"mau-import-cau-hoi.docx")
